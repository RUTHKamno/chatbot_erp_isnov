"""Artifacts generator — PDF, charts, PowerPoint, guide pipelines."""

from __future__ import annotations

import json
import logging
import time
import uuid
from pathlib import Path
from typing import Any

from app.agents.reformulator import assistant_disclaimer
from app.config import get_settings
from app.core.embeddings import embed_query
from app.graph.state import SMAState
from app.ingestion.store import get_store

logger = logging.getLogger(__name__)


def _detect_artifact_type(query: str) -> str:
    q = query.lower()
    if any(x in q for x in ("powerpoint", "pptx", "présentation", "presentation", "réunion", "reunion")):
        return "pptx"
    if any(x in q for x in ("pdf", "rapport", "jasper")):
        return "pdf"
    if any(x in q for x in ("graph", "visual", "chart", "courbe")):
        return "chart"
    if any(x in q for x in ("pipeline", "où cliquer", "ou cliquer", "comment", "étapes", "etapes")):
        return "guide_pipeline"
    return "text"


def _write_text_artifact(out_dir: Path, title: str, body: str) -> Path:
    path = out_dir / f"{title}-{uuid.uuid4().hex[:8]}.txt"
    path.write_text(body, encoding="utf-8")
    return path


def _write_guide_pipeline(out_dir: Path, query: str, docs: list[dict[str, Any]]) -> Path:
    steps = []
    if docs:
        for i, doc in enumerate(docs, start=1):
            steps.append(f"{i}. {doc.get('document', '')[:300]}")
    else:
        steps = [
            "1. Ouvrir le menu principal de l'ERP.",
            "2. Naviguer vers le module concerné par votre demande.",
            "3. Sélectionner l'action (consulter / modifier / valider).",
            "4. Enregistrer et vérifier le message de confirmation.",
        ]
    body = (
        f"Pipeline guide utilisateur\n"
        f"Question : {query}\n\n"
        + "\n".join(steps)
        + f"\n\n{assistant_disclaimer()}\n"
    )
    return _write_text_artifact(out_dir, "guide-pipeline", body)


def _write_pptx_stub(out_dir: Path, title: str, bullets: list[str]) -> Path:
    """Create a minimal PPTX when python-pptx is available."""
    path = out_dir / f"briefing-{uuid.uuid4().hex[:8]}.pptx"
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title[:100]
        body = slide.placeholders[1].text_frame
        body.clear()
        for i, bullet in enumerate(bullets[: get_settings().artifacts_max_pptx_slides]):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = bullet
            p.level = 0
            p.font.size = Pt(18)
        # Disclaimer slide
        slide2 = prs.slides.add_slide(prs.slide_layouts[1])
        slide2.shapes.title.text = "Avertissement"
        slide2.placeholders[1].text = assistant_disclaimer()
        prs.save(str(path))
        return path
    except Exception as exc:  # noqa: BLE001
        logger.warning("PPTX generation fallback: %s", exc)
        return _write_text_artifact(
            out_dir,
            "briefing",
            title + "\n\n" + "\n".join(f"- {b}" for b in bullets) + f"\n\n{assistant_disclaimer()}",
        )


def _write_pdf_stub(out_dir: Path, title: str, body: str) -> Path:
    path = out_dir / f"rapport-{uuid.uuid4().hex[:8]}.pdf"
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        c = canvas.Canvas(str(path), pagesize=A4)
        width, height = A4
        y = height - 50
        c.setFont("Helvetica-Bold", 14)
        c.drawString(50, y, title[:90])
        y -= 30
        c.setFont("Helvetica", 11)
        for line in body.splitlines():
            if y < 50:
                c.showPage()
                y = height - 50
                c.setFont("Helvetica", 11)
            c.drawString(50, y, line[:100])
            y -= 16
        c.save()
        return path
    except Exception as exc:  # noqa: BLE001
        logger.warning("PDF generation fallback: %s", exc)
        return _write_text_artifact(out_dir, "rapport", f"{title}\n\n{body}")


def _write_chart_stub(out_dir: Path, title: str) -> Path:
    path = out_dir / f"chart-{uuid.uuid4().hex[:8]}.png"
    try:
        import matplotlib.pyplot as plt

        fig, ax = plt.subplots(figsize=(8, 4))
        ax.plot([1, 2, 3, 4], [10, 14, 12, 18], marker="o")
        ax.set_title(title[:80])
        ax.set_xlabel("Période")
        ax.set_ylabel("Valeur")
        fig.tight_layout()
        fig.savefig(path)
        plt.close(fig)
        return path
    except Exception as exc:  # noqa: BLE001
        logger.warning("Chart generation fallback: %s", exc)
        meta = out_dir / f"chart-{uuid.uuid4().hex[:8]}.json"
        meta.write_text(json.dumps({"title": title, "error": str(exc)}), encoding="utf-8")
        return meta


async def artifacts_node(state: SMAState) -> dict[str, Any]:
    """Produce user-facing artifacts (guide pipeline, PDF, PPTX, charts)."""
    started = time.perf_counter()
    settings = get_settings()
    query = state.get("raw_query") or ""
    intent = state.get("intent") or ""
    out_dir = settings.artifacts_path
    trace = list(state.get("debug_trace") or [])
    paths: list[str] = []

    artifact_type = state.get("artifact_type") or _detect_artifact_type(query)

    retrieved: list[dict[str, Any]] = list(state.get("retrieved_docs") or [])
    if artifact_type in {"guide_pipeline", "text"} or intent == "user_guide":
        if not retrieved:
            try:
                claims = state.get("jwt_claims") or {}
                tenant_id = claims.get("tenant_id")
                user_role = claims.get("role")
                if tenant_id and user_role:
                    store = get_store()
                    query_embedding = await embed_query(query)
                    retrieved = store.search_enterprises(
                        query_vector=query_embedding,
                        tenant_id=tenant_id,
                        user_role=user_role,
                        top_k=4,
                    )
            except Exception as exc:  # noqa: BLE001
                logger.warning("Vector search skipped: %s", exc)

    if artifact_type == "pptx":
        bullets = [
            "Synthèse de la demande utilisateur",
            (state.get("answer") or query)[:200],
            "Indicateurs clés (à compléter avec Wren/DWH)",
            assistant_disclaimer(),
        ]
        path = _write_pptx_stub(out_dir, "Briefing BI", bullets)
        paths.append(str(path))
        answer = f"Présentation générée : `{path.name}`\n\n_{assistant_disclaimer()}_"
    elif artifact_type == "pdf":
        body = (state.get("answer") or query) + f"\n\n{assistant_disclaimer()}"
        path = _write_pdf_stub(out_dir, "Rapport ERP/BI", body)
        paths.append(str(path))
        answer = f"Rapport PDF généré : `{path.name}`\n\n_{assistant_disclaimer()}_"
    elif artifact_type == "chart":
        path = _write_chart_stub(out_dir, "Visualisation analytique")
        paths.append(str(path))
        answer = f"Graphique généré : `{path.name}`\n\n_{assistant_disclaimer()}_"
    else:
        path = _write_guide_pipeline(out_dir, query, retrieved)
        paths.append(str(path))
        content = path.read_text(encoding="utf-8")
        answer = content

    latency = (time.perf_counter() - started) * 1000
    if settings.enable_debug_trace:
        trace.append(
            {
                "agent": "artifacts",
                "action": "generate",
                "detail": f"type={artifact_type} files={paths}",
                "latency_ms": round(latency, 2),
                "ok": True,
            }
        )

    return {
        "artifact_type": artifact_type,
        "artifact_paths": paths,
        "retrieved_docs": retrieved,
        "answer": answer,
        "sources": ["artifacts", "vector_store"] if retrieved else ["artifacts"],
        "debug_trace": trace,
    }


async def guide_node(state: SMAState) -> dict[str, Any]:
    """User-guide path: RAG over ERP knowledge then format as pipeline."""
    return await artifacts_node({**state, "artifact_type": "guide_pipeline", "intent": "user_guide"})
